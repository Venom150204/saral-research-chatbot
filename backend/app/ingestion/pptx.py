"""PPTX extraction with slide-, shape-, paragraph-, table-, and notes provenance."""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from backend.app.domain import (
    BoundingBox,
    ExtractedDocument,
    ExtractionWarning,
    SourceBlock,
    SourceLocation,
    VisualElement,
)
from backend.app.storage import AssetStore, BinaryAsset

from .errors import ExtractionQualityError, UnsupportedFormatError


@dataclass(frozen=True, slots=True)
class PptxExtractionLimits:
    max_bytes: int = 50 * 1024 * 1024
    max_uncompressed_bytes: int = 250 * 1024 * 1024
    max_archive_entries: int = 5_000
    max_slides: int = 300

    def __post_init__(self) -> None:
        if min(
            self.max_bytes,
            self.max_uncompressed_bytes,
            self.max_archive_entries,
            self.max_slides,
        ) < 1:
            raise ValueError("all PPTX extraction limits must be positive")


def extract_pptx_bytes(
    payload: bytes,
    *,
    filename: str,
    limits: PptxExtractionLimits | None = None,
    asset_store: AssetStore | None = None,
) -> ExtractedDocument:
    active_limits = limits or PptxExtractionLimits()
    safe_filename = Path(filename).name
    if Path(safe_filename).suffix.lower() != ".pptx":
        raise UnsupportedFormatError("PPTX extractor only accepts .pptx uploads")
    if not payload:
        raise UnsupportedFormatError("PPTX payload is empty")
    if len(payload) > active_limits.max_bytes:
        raise UnsupportedFormatError(
            f"deck is {len(payload)} bytes; limit is {active_limits.max_bytes} bytes"
        )

    _validate_archive(payload, active_limits)
    try:
        presentation = Presentation(BytesIO(payload))
    except (PackageNotFoundError, KeyError, ValueError, BadZipFile) as exc:
        raise UnsupportedFormatError(f"PPTX package cannot be read: {exc}") from exc

    slide_count = len(presentation.slides)
    if slide_count < 1:
        raise ExtractionQualityError("deck contains no slides")
    if slide_count > active_limits.max_slides:
        raise UnsupportedFormatError(
            f"deck has {slide_count} slides; limit is {active_limits.max_slides}"
        )

    content_hash = sha256(payload).hexdigest()
    blocks: list[SourceBlock] = []
    visuals: list[VisualElement] = []
    warnings: list[ExtractionWarning] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        parent_id = f"sld_{content_hash[:12]}_{slide_number:04d}"
        slide_block_start = len(blocks)
        slide_visual_start = len(visuals)
        title_shape = slide.shapes.title
        title_shape_id = title_shape.shape_id if title_shape is not None else None

        for shape in slide.shapes:
            _extract_shape(
                shape=shape,
                slide_number=slide_number,
                parent_id=parent_id,
                title_shape_id=title_shape_id,
                content_hash=content_hash,
                blocks=blocks,
                visuals=visuals,
                asset_store=asset_store,
            )

        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                for paragraph_index, paragraph in enumerate(notes_frame.paragraphs):
                    if paragraph.text.strip():
                        _append_block(
                            blocks=blocks,
                            content_hash=content_hash,
                            parent_id=parent_id,
                            text=paragraph.text,
                            block_type="speaker_note",
                            location=SourceLocation(
                                slide_number=slide_number,
                                paragraph_index=paragraph_index,
                            ),
                        )

        if len(blocks) == slide_block_start:
            warnings.append(
                ExtractionWarning(
                    code="slide_without_extractable_text",
                    message="Slide contains no extractable text or speaker notes.",
                    slide_number=slide_number,
                )
            )
        if len(visuals) > slide_visual_start:
            warnings.append(
                ExtractionWarning(
                    code="visual_meaning_not_interpreted",
                    message=(
                        "Visual elements were located but their semantic content was not "
                        "interpreted by the PPTX baseline."
                    ),
                    slide_number=slide_number,
                )
            )

    if not blocks:
        raise ExtractionQualityError(
            "deck yielded no extractable text; visual analysis is required"
        )

    return ExtractedDocument(
        document_id=f"doc_{content_hash[:16]}",
        document_version=f"sha256:{content_hash}",
        content_hash=content_hash,
        filename=safe_filename,
        media_type=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
        source_kind="slide_deck",
        parser_name="python-pptx",
        parser_version=pptx.__version__,
        unit_count=slide_count,
        blocks=tuple(blocks),
        visuals=tuple(visuals),
        warnings=tuple(warnings),
    )


def _validate_archive(payload: bytes, limits: PptxExtractionLimits) -> None:
    try:
        with ZipFile(BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_archive_entries:
                raise UnsupportedFormatError(
                    f"PPTX has {len(entries)} archive entries; limit is "
                    f"{limits.max_archive_entries}"
                )
            uncompressed_size = sum(entry.file_size for entry in entries)
            if uncompressed_size > limits.max_uncompressed_bytes:
                raise UnsupportedFormatError(
                    f"PPTX expands to {uncompressed_size} bytes; limit is "
                    f"{limits.max_uncompressed_bytes} bytes"
                )
    except BadZipFile as exc:
        raise UnsupportedFormatError("PPTX is not a valid ZIP package") from exc


def _extract_shape(
    *,
    shape: object,
    slide_number: int,
    parent_id: str,
    title_shape_id: int | None,
    content_hash: str,
    blocks: list[SourceBlock],
    visuals: list[VisualElement],
    asset_store: AssetStore | None,
) -> None:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            _extract_shape(
                shape=child,
                slide_number=slide_number,
                parent_id=parent_id,
                title_shape_id=title_shape_id,
                content_hash=content_hash,
                blocks=blocks,
                visuals=visuals,
                asset_store=asset_store,
            )
        return

    box = BoundingBox(
        left=int(shape.left),
        top=int(shape.top),
        width=max(0, int(shape.width)),
        height=max(0, int(shape.height)),
    )

    if getattr(shape, "has_table", False):
        table = shape.table
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                if getattr(cell, "is_spanned", False) or not cell.text.strip():
                    continue
                _append_block(
                    blocks=blocks,
                    content_hash=content_hash,
                    parent_id=parent_id,
                    text=cell.text,
                    block_type="table_cell",
                    location=SourceLocation(
                        slide_number=slide_number,
                        shape_id=shape.shape_id,
                        table_row=row_index,
                        table_column=column_index,
                        bounding_box=box,
                    ),
                )

    if getattr(shape, "has_text_frame", False):
        for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
            if not paragraph.text.strip():
                continue
            block_type = "title" if shape.shape_id == title_shape_id else "slide_text"
            _append_block(
                blocks=blocks,
                content_hash=content_hash,
                parent_id=parent_id,
                text=paragraph.text,
                block_type=block_type,
                hierarchy_level=paragraph.level,
                location=SourceLocation(
                    slide_number=slide_number,
                    shape_id=shape.shape_id,
                    paragraph_index=paragraph_index,
                    bounding_box=box,
                ),
            )

    if getattr(shape, "has_chart", False):
        chart_markdown = _chart_to_markdown(shape.chart)
        if chart_markdown is not None:
            _append_block(
                blocks=blocks,
                content_hash=content_hash,
                parent_id=parent_id,
                text=chart_markdown,
                block_type="chart_data",
                location=SourceLocation(
                    slide_number=slide_number,
                    shape_id=shape.shape_id,
                    bounding_box=box,
                ),
            )

    is_visual = shape.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(
        shape, "has_chart", False
    )
    if is_visual:
        visual_type = "chart" if getattr(shape, "has_chart", False) else "picture"
        identity = (
            f"{content_hash}\0{slide_number}\0{shape.shape_id}\0{visual_type}"
        ).encode("utf-8")
        asset_hash: str | None = None
        media_type: str | None = None
        byte_size: int | None = None
        pixel_width: int | None = None
        pixel_height: int | None = None
        asset_uri: str | None = None
        if visual_type == "picture":
            image = shape.image
            image_bytes = image.blob
            asset_hash = sha256(image_bytes).hexdigest()
            media_type = image.content_type
            byte_size = len(image_bytes)
            pixel_width, pixel_height = image.size
            if asset_store is not None:
                asset_uri = asset_store.put(
                    BinaryAsset(
                        asset_hash=asset_hash,
                        media_type=media_type,
                        data=image_bytes,
                        original_filename=image.filename,
                    )
                )

        visuals.append(
            VisualElement(
                visual_id=f"vis_{sha256(identity).hexdigest()[:20]}",
                parent_id=parent_id,
                slide_number=slide_number,
                shape_id=shape.shape_id,
                shape_name=shape.name,
                visual_type=visual_type,
                bounding_box=box,
                asset_hash=asset_hash,
                media_type=media_type,
                byte_size=byte_size,
                pixel_width=pixel_width,
                pixel_height=pixel_height,
                asset_uri=asset_uri,
            )
        )


def _chart_to_markdown(chart: object) -> str | None:
    """Serialize native chart data without inferring a visual interpretation."""

    sections: list[str] = []
    if chart.has_title:
        title = chart.chart_title.text_frame.text.strip()
        if title:
            sections.append(f"Chart title: {title}")

    for plot in chart.plots:
        try:
            categories = [str(category) for category in plot.categories]
        except (AttributeError, TypeError, ValueError):
            categories = []
        series = list(plot.series)
        if not categories or not series:
            continue

        headers = ["Category", *(str(item.name) for item in series)]
        rows = [
            [
                category,
                *(
                    _format_chart_value(item.values[index])
                    if index < len(item.values)
                    else ""
                    for item in series
                ),
            ]
            for index, category in enumerate(categories)
        ]
        sections.append(_markdown_table(headers, rows))

    return "\n\n".join(sections) if sections else None


def _markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    def escape(value: str) -> str:
        return value.replace("|", "\\|").replace("\n", "<br>")

    header = "| " + " | ".join(escape(value) for value in headers) + " |"
    separator = "| " + " | ".join("---" for _ in headers) + " |"
    body = [
        "| " + " | ".join(escape(value) for value in row) + " |" for row in rows
    ]
    return "\n".join([header, separator, *body])


def _format_chart_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value)


def _append_block(
    *,
    blocks: list[SourceBlock],
    content_hash: str,
    parent_id: str,
    text: str,
    block_type: str,
    location: SourceLocation,
    hierarchy_level: int | None = None,
) -> None:
    ordinal = len(blocks)
    identity = (
        f"{content_hash}\0{parent_id}\0{ordinal}\0{block_type}\0{text}"
    ).encode("utf-8")
    blocks.append(
        SourceBlock(
            block_id=f"blk_{sha256(identity).hexdigest()[:20]}",
            parent_id=parent_id,
            ordinal=ordinal,
            text=text,
            block_type=block_type,
            hierarchy_level=hierarchy_level,
            location=location,
        )
    )
